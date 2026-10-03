"""Gera uma nova versão do TCC preservando o modelo Word original."""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
REFERENCE = ROOT / "TCC_ResNet_Completo_Template.docx"
DEFAULT_OUTPUT = ROOT / "TCC_ResNet_Rascunho_Experimental.docx"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--page-count", type=int, default=0)
    return parser.parse_args()


def set_run_font(run, name: str = "Times New Roman", size: float = 12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_paragraph(paragraph, text: str, *, bold=False, italic=False, size=12):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return paragraph


def remove_body_from(document: Document, start_index: int):
    paragraphs = document.paragraphs
    for paragraph in paragraphs[start_index:]:
        parent = paragraph._element.getparent()
        parent.remove(paragraph._element)


def add_body(document: Document, text: str, *, first_line=True, space_after=6):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Inches(0.49) if first_line else None
    p.paragraph_format.keep_together = False
    run = p.add_run(text)
    set_run_font(run, size=12)
    return p


def add_chapter(document: Document, number: int, title: str, *, first=False):
    p = document.add_paragraph()
    if not first:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(f"{number} {title.upper()}")
    set_run_font(run, size=14, bold=True)
    return p


def add_section_heading(document: Document, number: str, title: str):
    p = document.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(f"{number} {title}")
    set_run_font(run, size=12, bold=True)
    # Um separador mínimo e não vazio evita que o LibreOffice compacte o título
    # com a primeira linha do parágrafo seguinte quando há tabelas próximas.
    spacer = document.add_paragraph()
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer_run = spacer.add_run("\u00a0")
    set_run_font(spacer_run, size=1)
    return p


def add_bullet(document: Document, text: str):
    # O modelo não possui definições de numeração. Mantemos cada item como um
    # parágrafo recuado, sem simular marcadores com caracteres Unicode.
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.49)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text[0].upper() + text[1:])
    set_run_font(run, size=12)
    return p


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def add_table(document: Document, headers, rows, widths, caption: str):
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(4)
    run = cap.add_run(caption)
    set_run_font(run, size=10, bold=True)

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for i, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        set_cell_width(cell, width)
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(header))
        set_run_font(run, size=9, bold=True)
    set_repeat_table_header(table.rows[0])

    for row_values in rows:
        row = table.add_row()
        for i, (cell, value, width) in enumerate(zip(row.cells, row_values, widths)):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_run_font(run, size=8.5)

    source = document.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_before = Pt(3)
    source.paragraph_format.space_after = Pt(8)
    run = source.add_run("Fonte: elaborado pelo autor.")
    set_run_font(run, size=9)
    return table


def add_figure(document: Document, path: Path, caption: str, width=5.8):
    if not path.is_file():
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))

    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_together = True
    cap.paragraph_format.space_after = Pt(2)
    run = cap.add_run(caption)
    set_run_font(run, size=10, bold=True)

    source = document.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_after = Pt(8)
    run = source.add_run("Fonte: elaborado pelo autor a partir do conjunto EuroSAT.")
    set_run_font(run, size=9)


def build(output: Path, page_count: int):
    document = Document(REFERENCE)
    original_paragraphs = document.paragraphs

    # Remove exemplos pessoais e instruções genéricas herdados do modelo.
    for index in (113, 114):
        replace_paragraph(original_paragraphs[index], "", size=12)

    pages = str(page_count) if page_count else "a definir"
    replace_paragraph(
        original_paragraphs[146],
        f"ANJOS, E. J. Classificação de imagens com ResNet: uso de redes neurais convolucionais para classificação automatizada. 2026. {pages} f. Trabalho de Conclusão de Curso (MBA em Inteligência Artificial e Big Data) - Instituto de Ciências Matemáticas e de Computação, Universidade de São Paulo, São Carlos, 2026.",
        size=12,
    )
    replace_paragraph(
        original_paragraphs[148],
        "Este trabalho investiga o compromisso entre desempenho preditivo e custo computacional de redes residuais aplicadas à classificação de uso e cobertura do solo em imagens de satélite. Foram comparadas as arquiteturas ResNet18, ResNet34 e ResNet50 no conjunto EuroSAT RGB, composto por 27.000 imagens distribuídas em dez classes. Os experimentos utilizaram uma divisão estratificada e fixa de 70% para treinamento, 15% para validação e 15% para teste, transferência de aprendizado a partir do ImageNet e três sementes de treinamento. Além da acurácia e do macro-F1, foram medidos latência, throughput, memória de inferência e tamanho do checkpoint. A ResNet50 obteve o maior macro-F1 médio na validação, 98,8644%, enquanto a ResNet18 alcançou 98,5113%, com latência aproximadamente 2,56 vezes menor em inferência unitária. Aplicando uma regra de seleção registrada antes do teste, a ResNet18 foi escolhida como melhor compromisso de custo-benefício. No teste final, realizado uma única vez em 4.050 imagens, o modelo atingiu 98,6667% de acurácia e 98,6122% de macro-F1. Os resultados indicam que a ResNet18 oferece desempenho elevado com menor custo computacional no cenário avaliado, embora a generalização para outros sensores, regiões e condições temporais exija novos estudos.",
        size=12,
    )
    replace_paragraph(
        original_paragraphs[150],
        "Palavras-chave: classificação de imagens. ResNet. EuroSAT. sensoriamento remoto. transferência de aprendizado. custo-benefício.",
        size=12,
    )
    replace_paragraph(
        original_paragraphs[154],
        f"ANJOS, E. J. Image classification with ResNet: use of convolutional neural networks for automated classification. 2026. {pages} p. Course Completion Work (MBA in Artificial Intelligence and Big Data) - Institute of Mathematics and Computer Sciences, University of São Paulo, São Carlos, 2026.",
        size=12,
    )
    replace_paragraph(
        original_paragraphs[156],
        "This work investigates the trade-off between predictive performance and computational cost of residual networks applied to land-use and land-cover classification in satellite images. ResNet18, ResNet34 and ResNet50 were compared on the RGB EuroSAT dataset, which contains 27,000 images in ten classes. The experiments used a fixed stratified split of 70% for training, 15% for validation and 15% for testing, ImageNet transfer learning and three training seeds. Accuracy and macro-F1 were analyzed together with latency, throughput, inference memory and checkpoint size. ResNet50 achieved the best mean validation macro-F1, 98.8644%, while ResNet18 reached 98.5113% with approximately 2.56 times lower single-image latency. According to a selection rule registered before opening the test set, ResNet18 was chosen as the best cost-benefit solution. In the single final evaluation on 4,050 test images, it achieved 98.6667% accuracy and 98.6122% macro-F1. The results indicate that ResNet18 offers high predictive performance with lower computational cost in the evaluated setting, although generalization to other sensors, regions and temporal conditions requires further investigation.",
        size=12,
    )
    replace_paragraph(
        original_paragraphs[159],
        "Keywords: image classification. ResNet. EuroSAT. remote sensing. transfer learning. cost-benefit analysis.",
        size=12,
    )

    toc_lines = [
        "",
        "1 INTRODUÇÃO.............................................................................................................. 1",
        "2 FUNDAMENTAÇÃO TEÓRICA................................................................................... 3",
        "3 METODOLOGIA........................................................................................................... 6",
        "4 RESULTADOS E DISCUSSÃO..................................................................................... 8",
        "5 CONCLUSÃO............................................................................................................... 13",
        "REFERÊNCIAS................................................................................................................ 14",
    ]
    for offset, index in enumerate(range(164, 194)):
        text = toc_lines[offset] if offset < len(toc_lines) else ""
        replace_paragraph(original_paragraphs[index], text, bold=1 <= offset <= 6, size=12)

    remove_body_from(document, 194)

    # A epígrafe do modelo continha apenas instruções e autoria fictícia.
    # Remove-se o bloco opcional inteiro para não deixar uma página residual.
    for paragraph in document.paragraphs[117:145]:
        parent = paragraph._element.getparent()
        parent.remove(paragraph._element)

    last_sect_pr = document.sections[-1]._sectPr
    pg_num = last_sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        last_sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), "1")

    add_chapter(document, 1, "Introdução", first=True)
    add_body(document, "O uso de imagens orbitais ampliou a capacidade de observar áreas extensas e acompanhar transformações no território. Entretanto, o volume de dados produzido por programas de observação da Terra torna inviável a inspeção exclusivamente manual. Métodos de visão computacional podem apoiar esse processo ao atribuir automaticamente uma categoria de uso ou cobertura do solo a cada imagem, fornecendo informação útil para triagem, mapeamento e monitoramento.")
    add_body(document, "As redes neurais convolucionais passaram a ocupar posição central na classificação de imagens por aprenderem representações diretamente dos pixels. Entre essas redes, a Residual Network (ResNet) introduziu conexões de atalho que favorecem a otimização de modelos profundos (HE et al., 2016). Diferentes profundidades, porém, oferecem combinações distintas de desempenho, latência, memória e tamanho de armazenamento. Em uma aplicação real, o modelo com maior acurácia nem sempre representa a escolha mais adequada quando o ganho preditivo é pequeno e o custo computacional cresce de forma relevante.")
    add_body(document, "Este trabalho delimita o problema à classificação de imagens RGB do EuroSAT em dez classes: cultivo anual, floresta, vegetação herbácea, rodovia, área industrial, pastagem, cultivo permanente, área residencial, rio e mar/lago. O conjunto foi derivado de imagens Sentinel-2 e criado para pesquisas de uso e cobertura do solo (HELBER et al., 2019).")

    add_section_heading(document, "1.1", "Problema e questão de pesquisa")
    add_body(document, "O problema investigado consiste em selecionar uma arquitetura ResNet capaz de manter elevado desempenho de classificação com custo computacional compatível com uma futura aplicação de análise de imagens de satélite. A questão de pesquisa é: entre ResNet18, ResNet34 e ResNet50, qual arquitetura apresenta o melhor compromisso entre macro-F1 e custo de inferência no EuroSAT RGB, sob um protocolo experimental controlado?")

    add_section_heading(document, "1.2", "Objetivos")
    add_body(document, "O objetivo geral é comparar o desempenho preditivo e o custo computacional de três arquiteturas ResNet na classificação de uso e cobertura do solo em imagens do EuroSAT.")
    for item in [
        "implementar um pipeline reproduzível de treinamento, validação e teste em PyTorch;",
        "comparar ResNet18, ResNet34 e ResNet50 usando a mesma divisão de dados e os mesmos hiperparâmetros principais;",
        "avaliar o efeito da transferência de aprendizado e do congelamento do extrator de características;",
        "medir acurácia, macro-F1, latência, throughput, memória de inferência e tamanho do checkpoint;",
        "definir uma regra de escolha antes da abertura do conjunto de teste;",
        "analisar o desempenho final por classe e as principais confusões do modelo selecionado.",
    ]:
        add_bullet(document, item)

    add_section_heading(document, "1.3", "Organização do trabalho")
    add_body(document, "Além desta introdução, o Capítulo 2 apresenta os conceitos de classificação de imagens, redes convolucionais, aprendizado residual, ImageNet, transferência de aprendizado e EuroSAT. O Capítulo 3 descreve o protocolo experimental. O Capítulo 4 apresenta e discute os resultados. O Capítulo 5 reúne as conclusões, limitações e possibilidades de continuidade.")

    add_chapter(document, 2, "Fundamentação teórica")
    add_section_heading(document, "2.1", "Classificação de imagens e redes convolucionais")
    add_body(document, "Classificação de imagens é a tarefa de associar uma imagem a uma categoria previamente definida. Em redes convolucionais, filtros aprendidos identificam padrões locais, como bordas, formas e texturas. À medida que a informação atravessa as camadas, essas características são combinadas em representações mais abstratas. Ao final, uma camada de classificação produz escores para cada classe.")
    add_body(document, "Em sensoriamento remoto, a aparência das classes é influenciada por resolução espacial, sazonalidade, condições atmosféricas e características regionais. Categorias agrícolas podem compartilhar cores e texturas, enquanto elementos como floresta e corpos d'água tendem a apresentar padrões mais distintos. Por esse motivo, a análise por classe complementa a acurácia global.")

    add_section_heading(document, "2.2", "Redes residuais")
    add_body(document, "O aumento da profundidade de uma rede não garante melhoria automática e pode dificultar a otimização. A ResNet reformula blocos de camadas para aprender uma função residual F(x), cuja saída é somada à entrada x, resultando em H(x) = F(x) + x. As conexões de identidade oferecem um caminho direto para a propagação da informação e dos gradientes, facilitando o treinamento de redes profundas (HE et al., 2016).")
    add_body(document, "As variantes ResNet18 e ResNet34 utilizam blocos residuais básicos, enquanto a ResNet50 emprega blocos bottleneck. Assim, profundidade e organização dos blocos alteram número de parâmetros, custo do forward e capacidade de representação. Neste trabalho, essas diferenças são examinadas sob um mesmo conjunto de dados e protocolo.")

    add_section_heading(document, "2.3", "ImageNet e transferência de aprendizado")
    add_body(document, "O ImageNet foi proposto como uma base hierárquica de imagens em larga escala para apoiar pesquisas de reconhecimento visual (DENG et al., 2009). Pesos aprendidos nesse conjunto podem servir como inicialização para outra tarefa. Essa estratégia, denominada transferência de aprendizado, aproveita representações visuais gerais e pode reduzir o esforço necessário para aprender a partir de um conjunto-alvo menor.")
    add_body(document, "Foram avaliadas três situações: ajuste completo de pesos pré-treinados, treinamento da ResNet18 a partir de inicialização aleatória e uso da ResNet18 pré-treinada com backbone congelado. Essa comparação permite distinguir o efeito da transferência de aprendizado do efeito da arquitetura isoladamente.")

    add_section_heading(document, "2.4", "EuroSAT")
    add_body(document, "O EuroSAT reúne 27.000 imagens rotuladas em dez classes e foi construído a partir de dados do satélite Sentinel-2 (HELBER et al., 2019). A versão RGB foi adotada por ser diretamente compatível com modelos pré-treinados no ImageNet. A escolha também torna o experimento mais simples de reproduzir, embora descarte bandas espectrais que poderiam fornecer informação adicional.")
    add_figure(document, ROOT / "outputs/dataset_samples/eurosat_train_samples.png", "Figura 1 - Exemplos das dez classes do EuroSAT utilizados no estudo", width=5.7)

    add_section_heading(document, "2.5", "Métricas de desempenho e eficiência")
    add_body(document, "A acurácia corresponde à proporção total de classificações corretas. Precisão, recall e F1 foram calculados por classe. O macro-F1 é a média não ponderada do F1 das dez classes e atribui a mesma importância a cada categoria, independentemente do número de exemplos. Para avaliar eficiência, foram consideradas latência por imagem, throughput, pico de memória do forward e tamanho do checkpoint.")

    add_chapter(document, 3, "Metodologia")
    add_section_heading(document, "3.1", "Natureza da pesquisa")
    add_body(document, "A pesquisa é aplicada, quantitativa e experimental. O fator principal é a arquitetura da rede, com níveis ResNet18, ResNet34 e ResNet50. Também foram realizadas ablações sobre pré-treinamento e congelamento do backbone. As variáveis de resposta são métricas preditivas e medidas de custo computacional.")

    add_section_heading(document, "3.2", "Divisão e preparação dos dados")
    add_body(document, "A divisão foi estratificada por classe e gerada uma única vez com seed 42. Os índices foram persistidos em arquivo JSON e reutilizados em todos os experimentos. Dessa forma, variações de resultado não decorrem de mudanças nas imagens de treinamento e validação. O conjunto de teste permaneceu sem consulta durante treinamento, comparação, benchmark e seleção do modelo.")
    add_table(
        document,
        ["Partição", "Proporção", "Imagens", "Finalidade"],
        [
            ["Treino", "70%", "18.900", "Ajuste dos parâmetros"],
            ["Validação", "15%", "4.050", "Comparação e seleção"],
            ["Teste", "15%", "4.050", "Avaliação final única"],
        ],
        [1900, 1500, 1600, 4000],
        "Tabela 1 - Divisão estratificada do EuroSAT",
    )
    add_body(document, "No treinamento, foram aplicados recorte aleatório redimensionado para 224 x 224 pixels, inversões horizontal e vertical, rotação aleatória de até 20 graus e normalização com média e desvio-padrão do ImageNet. Na validação e no teste, as imagens foram redimensionadas para 256 pixels, submetidas a recorte central de 224 x 224 e normalizadas da mesma forma. Apenas o treino recebeu transformações aleatórias.")

    add_section_heading(document, "3.3", "Configuração de treinamento")
    add_body(document, "A implementação utilizou Python, PyTorch e torchvision. PyTorch oferece execução eficiente em aceleradores e uma interface imperativa adequada à experimentação (PASZKE et al., 2019). Os treinamentos foram realizados em uma GPU NVIDIA GeForce RTX 3060 com 12 GB. A última camada totalmente conectada foi substituída para produzir dez saídas.")
    add_table(
        document,
        ["Parâmetro", "Valor", "Parâmetro", "Valor"],
        [
            ["Épocas", "15", "Batch size", "64"],
            ["Taxa inicial", "0,0001", "Weight decay", "0,0001"],
            ["Otimizador", "AdamW", "Função de perda", "Cross-entropy"],
            ["Scheduler", "ReduceLROnPlateau", "Paciência", "2 épocas"],
            ["Seeds de treino", "42, 123 e 2026", "Seed da divisão", "42"],
            ["Precisão", "FP32", "Dispositivo", "CUDA"],
        ],
        [2200, 2300, 2200, 2300],
        "Tabela 2 - Configuração dos experimentos principais",
    )
    add_body(document, "O checkpoint de cada execução foi determinado pela menor perda de validação. A seed de treinamento controlou inicialização e aleatoriedade das transformações, enquanto a seed da divisão permaneceu fixa. Para cada arquitetura principal, foram executadas três seeds, permitindo reportar média e desvio-padrão amostral.")

    add_section_heading(document, "3.4", "Baselines e ablações")
    add_body(document, "A baseline ingênua sempre prevê a classe mais frequente no treino. Embora não utilize características das imagens, ela estabelece um limite inferior transparente. A ResNet18 treinada do zero mede o desempenho sem conhecimento transferido. As execuções com backbone congelado treinam apenas a camada final e avaliam uma alternativa de menor custo de ajuste.")

    add_section_heading(document, "3.5", "Benchmark de inferência")
    add_body(document, "O benchmark utilizou os checkpoints da seed 42 e exclusivamente a validação. Foram medidos batch sizes 1 e 64, com 20 lotes de aquecimento, cinco repetições e 100 forwards sincronizados por repetição na medição exclusiva do modelo. O tempo de forward foi obtido com eventos CUDA. O pipeline completo incluiu leitura, transformações, transferência para a GPU e execução do modelo. A memória principal de comparação correspondeu ao pico do forward sincronizado, reduzindo interferências do enfileiramento assíncrono.")

    add_section_heading(document, "3.6", "Regra de seleção e uso do teste")
    add_body(document, "Antes de acessar o teste, definiu-se que seriam elegíveis arquiteturas até 0,5 ponto percentual abaixo do melhor macro-F1 médio na validação. Entre as elegíveis, seria selecionada a menor latência média do modelo com batch size 1. O checkpoint representativo seria o da seed 42. A decisão foi registrada em arquivo antes da avaliação final. O teste foi executado uma única vez e não foi usado para alterar arquitetura, seed ou hiperparâmetros.")

    add_chapter(document, 4, "Resultados e discussão")
    add_section_heading(document, "4.1", "Baselines e efeito da transferência de aprendizado")
    add_table(
        document,
        ["Configuração", "Acurácia", "Macro-F1", "Interpretação"],
        [
            ["Baseline majoritária", "11,11%", "2,00%", "Limite inferior"],
            ["ResNet18 congelada, lr=10⁻⁴", "81,60%", "81,47%", "Ajuste insuficiente"],
            ["ResNet18 congelada, lr=10⁻³", "90,25%", "90,08%", "Melhora com taxa maior"],
            ["ResNet18 do zero", "96,25%", "96,13%", "Aprendizado sem ImageNet"],
            ["ResNet18 pré-treinada", "98,49%", "98,46%", "Seed 42; ajuste completo"],
        ],
        [3100, 1600, 1600, 2700],
        "Tabela 3 - Baselines e ablações na validação",
    )
    add_body(document, "Todas as redes superaram amplamente a baseline majoritária, demonstrando que as imagens contêm padrões discriminativos aprendidos pelos modelos. A ResNet18 treinada do zero atingiu macro-F1 de 96,13%, mas permaneceu aproximadamente 2,32 pontos percentuais abaixo da versão pré-treinada da seed 42. Isso evidencia a contribuição da transferência de aprendizado no protocolo adotado.")
    add_body(document, "Congelar o backbone reduziu o desempenho, especialmente com taxa de aprendizado de 10⁻⁴. O aumento para 10⁻³ elevou o macro-F1 para 90,08%, mas ainda ficou distante do ajuste completo. Portanto, a economia de treinamento obtida pelo congelamento produziu perda preditiva relevante neste conjunto.")
    add_figure(document, ROOT / "outputs/baseline_resnet18/learning_curves.png", "Figura 2 - Curvas de aprendizado da ResNet18 pré-treinada, seed 42", width=5.8)

    add_section_heading(document, "4.2", "Comparação entre ResNet18, ResNet34 e ResNet50")
    add_table(
        document,
        ["Modelo", "Macro-F1 médio", "Latência B1", "Throughput B64", "Memória B64", "Checkpoint"],
        [
            ["ResNet18", "98,5113% ± 0,0936", "2,402 ± 0,045 ms", "1.938,76 img/s", "479,58 MB", "42,73 MB"],
            ["ResNet34", "98,5454% ± 0,0580", "4,478 ± 0,148 ms", "1.089,18 img/s", "518,18 MB", "81,35 MB"],
            ["ResNet50", "98,8644% ± 0,0738", "6,152 ± 0,079 ms", "624,58 img/s", "773,24 MB", "90,06 MB"],
        ],
        [1500, 1900, 1750, 1450, 1200, 1200],
        "Tabela 4 - Desempenho médio e custo computacional das arquiteturas",
    )
    add_body(document, "A ResNet50 obteve o melhor desempenho absoluto, com vantagem de aproximadamente 0,353 ponto percentual de macro-F1 sobre a ResNet18. Entretanto, sua latência com batch 1 foi cerca de 2,56 vezes maior e seu checkpoint foi aproximadamente 2,11 vezes maior. Em batch 64, a ResNet18 apresentou throughput exclusivo do modelo cerca de 3,10 vezes superior.")
    add_body(document, "A ResNet34 apresentou macro-F1 médio apenas 0,034 ponto percentual superior ao da ResNet18, mas com latência aproximadamente 1,86 vez maior. Assim, ela não ofereceu vantagem suficiente para compensar o custo adicional. As barras de desvio-padrão mostram baixa variação entre seeds nas três arquiteturas, o que sustenta a estabilidade da comparação no split utilizado.")
    add_figure(document, ROOT / "outputs/plots/pareto_overview.png", "Figura 3 - Relações de Pareto entre macro-F1 e métricas de custo", width=5.9)

    add_section_heading(document, "4.3", "Seleção do modelo")
    add_body(document, "As três arquiteturas ficaram dentro da margem de 0,5 ponto percentual definida previamente. A segunda etapa da regra selecionou a ResNet18 por apresentar a menor latência em batch 1. Essa decisão não afirma que a ResNet18 é universalmente superior; ela expressa uma preferência explícita por eficiência quando o desempenho permanece próximo do melhor valor observado.")

    add_section_heading(document, "4.4", "Avaliação final no teste")
    add_body(document, "O checkpoint da ResNet18 pré-treinada, seed 42, foi avaliado uma única vez nas 4.050 imagens de teste. A acurácia foi 98,6667% e o macro-F1 foi 98,6122%. Em relação à validação do mesmo checkpoint, as diferenças foram +0,173 ponto percentual em acurácia e +0,155 ponto percentual em macro-F1. A proximidade entre os resultados indica boa generalização dentro da distribuição do EuroSAT.")
    add_figure(document, ROOT / "outputs/final_resnet18_test/validation_vs_test.png", "Figura 4 - Comparação global entre validação e teste", width=5.4)
    add_figure(document, ROOT / "outputs/final_resnet18_test/test_confusion_matrix_normalized.png", "Figura 5 - Matriz de confusão normalizada no teste final", width=5.6)
    add_figure(document, ROOT / "outputs/final_resnet18_test/test_f1_per_class.png", "Figura 6 - F1-score por classe no teste final", width=5.7)

    add_section_heading(document, "4.5", "Análise por classe")
    add_table(
        document,
        ["Classe", "Precisão", "Recall", "F1", "Suporte"],
        [
            ["Cultivo anual", "97,37%", "98,89%", "98,13%", "450"],
            ["Floresta", "99,78%", "99,56%", "99,67%", "450"],
            ["Vegetação herbácea", "98,00%", "98,22%", "98,11%", "450"],
            ["Rodovia", "98,41%", "99,20%", "98,80%", "375"],
            ["Área industrial", "98,67%", "98,93%", "98,80%", "375"],
            ["Pastagem", "97,04%", "98,33%", "97,68%", "300"],
            ["Cultivo permanente", "98,63%", "96,00%", "97,30%", "375"],
            ["Área residencial", "99,33%", "99,33%", "99,33%", "450"],
            ["Rio", "99,46%", "97,60%", "98,52%", "375"],
            ["Mar/Lago", "99,56%", "100,00%", "99,78%", "450"],
        ],
        [3100, 1550, 1550, 1450, 1350],
        "Tabela 5 - Métricas por classe no teste final",
    )
    add_body(document, "Cultivo permanente apresentou o menor F1, 97,30%. As maiores confusões dirigidas foram oito imagens de cultivo permanente classificadas como cultivo anual e seis como vegetação herbácea. Esse padrão é coerente com a semelhança de texturas, cores e organização espacial entre áreas agrícolas. Também ocorreram quatro classificações de rio como rodovia, possivelmente associadas a formas lineares presentes nas duas classes.")
    add_body(document, "Mar/lago e floresta obtiveram os maiores valores de F1, respectivamente 99,78% e 99,67%. Essas categorias apresentam características visuais mais estáveis no conjunto avaliado. A ausência de uma classe com desempenho muito inferior às demais explica a proximidade entre acurácia e macro-F1.")

    add_section_heading(document, "4.6", "Limitações e ameaças à validade")
    for item in [
        "Os resultados são internos ao EuroSAT RGB e não demonstram generalização para outros satélites, resoluções, regiões geográficas ou períodos do ano.",
        "O uso de apenas três seeds reduz, mas não elimina, a incerteza associada ao treinamento estocástico.",
        "O benchmark foi realizado em uma única configuração de hardware; valores absolutos de latência e memória podem mudar em outros dispositivos.",
        "A comparação foi limitada à família ResNet e não inclui arquiteturas recentes, modelos compactos ou métodos multiespectrais.",
        "As imagens do EuroSAT são recortes independentes; o estudo não avalia mapas contínuos, detecção de mudanças ou segmentação semântica.",
    ]:
        add_bullet(document, item)

    add_chapter(document, 5, "Conclusão")
    add_body(document, "Este trabalho comparou ResNet18, ResNet34 e ResNet50 para classificação de uso e cobertura do solo no EuroSAT RGB. O protocolo utilizou divisão estratificada fixa, três seeds por arquitetura, transferência de aprendizado, métricas preditivas e um benchmark de inferência. A seleção do modelo foi registrada antes da avaliação no teste, reduzindo o risco de adaptar a decisão ao resultado final.")
    add_body(document, "A ResNet50 apresentou o maior macro-F1 médio na validação, 98,8644%, mas exigiu maior latência, memória e armazenamento. A ResNet18 atingiu 98,5113% de macro-F1 médio com latência de 2,402 ms por imagem em batch 1, tornando-se a escolhida pela regra de custo-benefício. No teste final, obteve 98,6667% de acurácia e 98,6122% de macro-F1.")
    add_body(document, "A questão de pesquisa pode ser respondida da seguinte forma: no cenário experimental avaliado, a ResNet18 pré-treinada e ajustada integralmente apresentou o melhor compromisso entre desempenho e custo. A vantagem preditiva da ResNet50 foi real, porém pequena diante do aumento de latência e tamanho. Para uma aplicação em que cada fração de acurácia seja prioritária e haja recursos disponíveis, a ResNet50 permanece uma alternativa; para processamento eficiente, a ResNet18 é a recomendação sustentada pelos dados.")
    add_body(document, "Como continuidade, recomenda-se testar validação geográfica e temporal, imagens provenientes de outras regiões, bandas multiespectrais, quantização, poda, arquiteturas compactas e avaliação em dispositivos de borda. Também seria relevante investigar explicabilidade e confiança das previsões, especialmente nas classes agrícolas que concentraram as principais confusões.")

    p = document.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("REFERÊNCIAS")
    set_run_font(run, size=14, bold=True)
    references = [
        "DENG, J. et al. ImageNet: a large-scale hierarchical image database. In: IEEE CONFERENCE ON COMPUTER VISION AND PATTERN RECOGNITION, 2009. Anais [...]. p. 248-255. DOI: 10.1109/CVPR.2009.5206848.",
        "HE, K.; ZHANG, X.; REN, S.; SUN, J. Deep residual learning for image recognition. In: IEEE CONFERENCE ON COMPUTER VISION AND PATTERN RECOGNITION, 2016. Anais [...]. p. 770-778.",
        "HELBER, P.; BISCHKE, B.; DENGEL, A.; BORTH, D. EuroSAT: a novel dataset and deep learning benchmark for land use and land cover classification. IEEE Journal of Selected Topics in Applied Earth Observations and Remote Sensing, v. 12, n. 7, p. 2217-2226, 2019. DOI: 10.1109/JSTARS.2019.2918242.",
        "PASZKE, A. et al. PyTorch: an imperative style, high-performance deep learning library. In: ADVANCES IN NEURAL INFORMATION PROCESSING SYSTEMS, 2019. Anais [...]. v. 32, p. 8024-8035.",
    ]
    for reference in references:
        p = add_body(document, reference, first_line=False, space_after=12)
        p.paragraph_format.left_indent = None

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    args = parse_args()
    if args.output.resolve() == REFERENCE.resolve():
        raise ValueError("O arquivo de saída não pode ser o modelo original.")
    build(args.output.resolve(), args.page_count)
