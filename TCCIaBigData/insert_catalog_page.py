from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE_DOCX = Path(r"D:\TCCIaBigData\TCC_ResNet_FINAL_revisado_ABNT_paginacao_sequencial.docx")
CATALOG_IMAGE = Path(r"D:\TCCIaBigData\tmp\ficha_render\ficha.png")
OUTPUT_DOCX = Path(r"D:\TCCIaBigData\TCC_ResNet_FINAL_revisado_ABNT_ficha_catalografica.docx")


def convert_inline_to_full_page_anchor(drawing, page_width, page_height):
    inline = drawing.find(qn("wp:inline"))
    if inline is None:
        raise RuntimeError("The inserted image was not created as an inline drawing")

    children = {child.tag: child for child in list(inline)}
    required = [
        qn("wp:extent"),
        qn("wp:docPr"),
        qn("wp:cNvGraphicFramePr"),
        qn("a:graphic"),
    ]
    missing = [tag for tag in required if tag not in children]
    if missing:
        raise RuntimeError(f"Missing drawing components: {missing}")

    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "1",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(name, value)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "page")
    pos_h = OxmlElement("wp:posOffset")
    pos_h.text = "0"
    position_h.append(pos_h)
    anchor.append(position_h)

    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "page")
    pos_v = OxmlElement("wp:posOffset")
    pos_v.text = "0"
    position_v.append(pos_v)
    anchor.append(position_v)

    extent = children[qn("wp:extent")]
    extent.set("cx", str(page_width))
    extent.set("cy", str(page_height))
    anchor.append(extent)

    effect_extent = children.get(qn("wp:effectExtent"))
    if effect_extent is not None:
        anchor.append(effect_extent)

    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(children[qn("wp:docPr")])
    anchor.append(children[qn("wp:cNvGraphicFramePr")])
    anchor.append(children[qn("a:graphic")])

    drawing.replace(inline, anchor)


def main():
    document = Document(SOURCE_DOCX)
    if len(document.paragraphs) < 35:
        raise RuntimeError("Unexpected document structure")

    # The third physical page is the intentionally blank page between the
    # title page and the abstract. It is bounded by page-break paragraphs 30
    # and 34 in this document. Anchor the PDF raster to paragraph 31 so it
    # cannot reflow or alter the following pagination.
    target = document.paragraphs[31]
    target.alignment = WD_ALIGN_PARAGRAPH.LEFT
    target.paragraph_format.space_before = 0
    target.paragraph_format.space_after = 0
    target.paragraph_format.line_spacing = 1

    for run in list(target.runs):
        target._p.remove(run._r)

    section = document.sections[1]
    run = target.add_run()
    run.add_picture(
        str(CATALOG_IMAGE),
        width=section.page_width,
        height=section.page_height,
    )
    drawing = run._r.find(qn("w:drawing"))
    if drawing is None:
        raise RuntimeError("Inserted image drawing not found")
    convert_inline_to_full_page_anchor(
        drawing,
        int(section.page_width),
        int(section.page_height),
    )

    document.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
